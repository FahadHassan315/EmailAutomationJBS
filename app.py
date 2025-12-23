import streamlit as st
import os
from docx import Document
import re
import urllib.parse
import base64
from datetime import datetime, timedelta

# Set page config
st.set_page_config(
    page_title="Email Template Generator",
    page_icon="📧",
    layout="centered"
)

TEMPLATE_FOLDER = "Templates"

def read_docx(file_path):
    """Read Word document and return full text"""
    try:
        doc = Document(file_path)
        full_text = []
        for para in doc.paragraphs:
            if para.text.strip():  # Only add non-empty paragraphs
                full_text.append(para.text)
        return "\n".join(full_text)
    except Exception as e:
        st.error(f"Error reading document: {e}")
        return None

def extract_subject_and_body(content):
    """Extract subject line and body from content"""
    lines = content.split('\n')
    subject = ""
    body_lines = []
    
    subject_found = False
    for line in lines:
        if line.strip().lower().startswith('subject:'):
            subject = line.split(':', 1)[1].strip()
            subject_found = True
        elif subject_found:
            body_lines.append(line)
        else:
            body_lines.append(line)
    
    body = '\n'.join(body_lines).strip()
    return subject, body

def extract_placeholders(text):
    """Extract unique placeholders in {PlaceholderName} format"""
    placeholders = re.findall(r'\{([^}]+)\}', text)
    # Preserve order but remove duplicates
    seen = set()
    unique_placeholders = []
    for ph in placeholders:
        if ph not in seen:
            seen.add(ph)
            unique_placeholders.append(ph)
    return unique_placeholders

def replace_placeholders(text, values):
    """Replace placeholders with values"""
    result = text
    for placeholder, value in values.items():
        result = result.replace(f"{{{placeholder}}}", value)
    return result

def create_outlook_web_link(subject, body, to="", cc="", bcc=""):
    """Create Outlook Web deep link"""
    # Outlook Web deep link only supports plain text
    # Use proper line break encoding for URLs
    body_encoded = body.replace('\n', '%0D%0A')
    
    # Build the Outlook Web compose URL
    params = {
        'subject': urllib.parse.quote(subject),
        'body': body_encoded
    }
    
    if to:
        params['to'] = urllib.parse.quote(to)
    if cc:
        params['cc'] = urllib.parse.quote(cc)
    if bcc:
        params['bcc'] = urllib.parse.quote(bcc)
    
    # Build query string
    query_string = '&'.join([f"{k}={v}" for k, v in params.items()])
    
    outlook_url = f"https://outlook.office.com/mail/deeplink/compose?{query_string}"
    
    return outlook_url

def create_calendar_meeting_link(subject, body, attendees, start_time, end_time, location=""):
    """Create Outlook Calendar deep link for meeting"""
    # Format times for Outlook Calendar (ISO 8601 format)
    start_iso = start_time.strftime('%Y-%m-%dT%H:%M:%S')
    end_iso = end_time.strftime('%Y-%m-%dT%H:%M:%S')
    
    # Convert body to HTML with proper formatting
    html_body = body.replace('\n\n', '</p><p>').replace('\n', '<br>')
    html_body = f'<p>{html_body}</p>'
    
    # Build the Outlook Calendar compose URL
    params = {
        'subject': urllib.parse.quote(subject),
        'body': urllib.parse.quote(html_body),
        'startdt': urllib.parse.quote(start_iso),
        'enddt': urllib.parse.quote(end_iso),
        'path': '/calendar/action/compose'
    }
    
    if attendees:
        params['to'] = urllib.parse.quote(attendees)
    
    if location:
        params['location'] = urllib.parse.quote(location)
    
    # Build query string
    query_string = '&'.join([f"{k}={v}" for k, v in params.items()])
    
    calendar_url = f"https://outlook.office.com/calendar/0/deeplink/compose?{query_string}"
    
    return calendar_url

# Main app
st.title("📧 Email Template Generator")
st.markdown("**Generate professional emails and calendar meetings in Outlook Web**")

# Check if Templates folder exists
if not os.path.exists(TEMPLATE_FOLDER):
    st.error(f"❌ '{TEMPLATE_FOLDER}' folder not found!")
    st.info("Please create a 'Templates' folder in the same directory as this app and add your .docx files.")
    st.stop()

# Load templates
template_files = [f for f in os.listdir(TEMPLATE_FOLDER) if f.endswith('.docx') and not f.startswith('~$')]

if not template_files:
    st.error("❌ No .docx templates found in the 'Templates' folder.")
    st.stop()

# Template selection
st.subheader("1️⃣ Select Template")
template_name = st.selectbox(
    "Choose an email template:",
    template_files,
    format_func=lambda x: x.replace('.docx', '').replace('_', ' ')
)

# Check if this is the Online Interview template
is_interview_template = 'Online_Interview' in template_name

# Load selected template
doc_path = os.path.join(TEMPLATE_FOLDER, template_name)
content = read_docx(doc_path)

if content is None:
    st.stop()

# Extract subject and body
subject_text, body_text = extract_subject_and_body(content)

# Find all placeholders
placeholders = extract_placeholders(content)

# Show template preview
with st.expander("📄 View Template", expanded=False):
    st.text(content)

# Input fields for placeholders
if placeholders:
    st.subheader("2️⃣ Fill in Details")
    values = {}
    
    # Create columns for better layout
    col1, col2 = st.columns(2)
    
    for idx, placeholder in enumerate(placeholders):
        # Alternate between columns
        with col1 if idx % 2 == 0 else col2:
            values[placeholder] = st.text_input(
                placeholder.replace('_', ' ').title(),
                key=placeholder,
                placeholder=f"Enter {placeholder.lower()}"
            )
    
    # Replace placeholders
    final_subject = replace_placeholders(subject_text, values)
    final_body = replace_placeholders(body_text, values)
else:
    final_subject = subject_text
    final_body = body_text
    st.info("ℹ️ This template has no placeholders.")

# Calendar Meeting Section (only for Online Interview template)
meeting_start_time = None
meeting_end_time = None
meeting_link = ""
meeting_attendees = ""

if is_interview_template:
    st.markdown("---")
    st.subheader("📅 Interview Meeting Details")
    st.info("💡 Fill in these details to create a calendar meeting invitation")
    
    col1, col2 = st.columns(2)
    
    with col1:
        meeting_date = st.date_input(
            "Interview Date",
            value=datetime.now() + timedelta(days=7),
            min_value=datetime.now().date()
        )
        meeting_start = st.time_input(
            "Start Time",
            value=datetime.strptime("10:00", "%H:%M").time()
        )
    
    with col2:
        duration = st.selectbox(
            "Duration",
            options=[30, 45, 60, 90, 120],
            index=2,
            format_func=lambda x: f"{x} minutes"
        )
        meeting_end = (datetime.combine(meeting_date, meeting_start) + timedelta(minutes=duration)).time()
        st.time_input(
            "End Time",
            value=meeting_end,
            disabled=True
        )
    
    # Combine date and time
    meeting_start_time = datetime.combine(meeting_date, meeting_start)
    meeting_end_time = datetime.combine(meeting_date, meeting_end)
    
    # Meeting platform
    meeting_platform = st.selectbox(
        "Meeting Platform",
        options=["Microsoft Teams", "Zoom", "Google Meet", "Other"],
        index=0
    )
    
    meeting_link = st.text_input(
        "Meeting Link",
        placeholder="https://teams.microsoft.com/l/meetup-join/...",
        help="Paste the meeting link from Teams/Zoom/Google Meet"
    )
    
    meeting_attendees = st.text_input(
        "Candidate Email",
        placeholder="candidate@example.com",
        help="Enter the candidate's email address"
    )
    
    # Add meeting details to email body
    if meeting_link:
        meeting_details = f"\n\n📅 Interview Details:\nDate: {meeting_date.strftime('%B %d, %Y')}\nTime: {meeting_start.strftime('%I:%M %p')} - {meeting_end.strftime('%I:%M %p')}\nPlatform: {meeting_platform}\nMeeting Link: {meeting_link}"
        final_body = final_body + meeting_details

# Optional recipient fields
st.subheader("3️⃣ Recipients (Optional)" if not is_interview_template else "3️⃣ Additional Recipients (Optional)")
col1, col2, col3 = st.columns(3)

with col1:
    to_email = st.text_input("To:", placeholder="recipient@example.com", value=meeting_attendees if is_interview_template else "")
with col2:
    cc_email = st.text_input("CC:", placeholder="cc@example.com")
with col3:
    bcc_email = st.text_input("BCC:", placeholder="bcc@example.com")

# Preview section
st.subheader("4️⃣ Preview")

# Show preview in a nice box
with st.container():
    st.markdown("**Subject:**")
    st.code(final_subject if final_subject else "No subject")
    
    st.markdown("**Email Body:**")
    st.text_area("", final_body, height=300, disabled=True, label_visibility="collapsed")

# Generate Outlook Web link
outlook_link = create_outlook_web_link(
    subject=final_subject,
    body=final_body,
    to=to_email,
    cc=cc_email,
    bcc=bcc_email
)

# Main action buttons
st.subheader("5️⃣ Send Email & Create Meeting")

# Create two columns for buttons
col1, col2 = st.columns(2)

with col1:
    # Email button
    st.markdown(f"""
    <a href="{outlook_link}" target="_blank" style="text-decoration: none;">
        <button style="
            background-color: #0078D4;
            color: white;
            padding: 12px 24px;
            font-size: 16px;
            font-weight: bold;
            border: none;
            border-radius: 4px;
            cursor: pointer;
            width: 100%;
            margin-bottom: 10px;
        ">
            📧 Open Email in Outlook
        </button>
    </a>
    """, unsafe_allow_html=True)

with col2:
    # Calendar button (only for interview template)
    if is_interview_template and meeting_start_time and meeting_attendees:
        calendar_link = create_calendar_meeting_link(
            subject=final_subject,
            body=final_body,
            attendees=meeting_attendees,
            start_time=meeting_start_time,
            end_time=meeting_end_time,
            location=meeting_link if meeting_link else meeting_platform
        )
        
        st.markdown(f"""
        <a href="{calendar_link}" target="_blank" style="text-decoration: none;">
            <button style="
                background-color: #28A745;
                color: white;
                padding: 12px 24px;
                font-size: 16px;
                font-weight: bold;
                border: none;
                border-radius: 4px;
                cursor: pointer;
                width: 100%;
                margin-bottom: 10px;
            ">
                📅 Create Calendar Meeting
            </button>
        </a>
        """, unsafe_allow_html=True)
    elif is_interview_template:
        st.markdown("""
        <button style="
            background-color: #CCCCCC;
            color: #666666;
            padding: 12px 24px;
            font-size: 16px;
            font-weight: bold;
            border: none;
            border-radius: 4px;
            cursor: not-allowed;
            width: 100%;
            margin-bottom: 10px;
        " disabled>
            📅 Create Calendar Meeting
        </button>
        """, unsafe_allow_html=True)
        st.caption("⚠️ Fill in meeting details and candidate email to enable")

if is_interview_template:
    st.info("💡 Click '📧 Open Email' to send the invitation, then '📅 Create Calendar Meeting' to send the calendar invite!")
else:
    st.info("💡 Click the button above to open this email directly in Outlook Web. The email will be pre-filled and ready to send!")

# Alternative copy options
st.markdown("---")
st.subheader("📋 Alternative: Copy & Paste")

col1, col2 = st.columns(2)

with col1:
    if st.button("📋 Copy Email Body", use_container_width=True):
        st.code(final_body, language=None)
        st.success("✅ Email body displayed above - select and copy (Ctrl+C)")

with col2:
    if st.button("📋 Copy Subject", use_container_width=True):
        st.code(final_subject, language=None)
        st.success("✅ Subject displayed above - select and copy (Ctrl+C)")

# Instructions
st.markdown("---")

# Footer
st.markdown("---")
st.caption("💼 JBS Email Template Generator | Made with Streamlit")